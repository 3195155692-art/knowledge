from __future__ import annotations

import copy
import os
import re
import shutil
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


BASE = Path("/Users/oldma/knowledge/artical/Brain-Computer Interface/.template_deck_work/组会_template_base.pptx")
PATCHED = Path("/Users/oldma/knowledge/artical/Brain-Computer Interface/.template_deck_work/组会_template_patched.pptx")
OUT = Path("/Users/oldma/Desktop/data/Project materials/review/manuscript/round25/组会_完整版.pptx")

FONT_CN = "Heiti SC"
FONT_EN = "Times New Roman"

INK = RGBColor(0x14, 0x1A, 0x22)
SUB = RGBColor(0x5B, 0x66, 0x72)
ACCENT = RGBColor(0x1C, 0x4F, 0x72)
ALERT = RGBColor(0x9B, 0x3F, 0x3F)
PALE = RGBColor(0xF4, 0xF6, 0xF8)
PALE_BLUE = RGBColor(0xEA, 0xF1, 0xF6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def patch_package_fonts(src: Path, dst: Path) -> None:
    replacements = {
        "PingFang SC Semibold": FONT_CN,
        "PingFang SC": FONT_CN,
        "苹方-简": FONT_CN,
    }
    with zipfile.ZipFile(src, "r") as zin, zipfile.ZipFile(dst, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.endswith(".xml") or item.filename.endswith(".rels"):
                text = data.decode("utf-8")
                for old, new in replacements.items():
                    text = text.replace(old, new)
                data = text.encode("utf-8")
            zout.writestr(item, data)


def set_run_font(run, size: float | None = None, bold: bool | None = None, name: str = FONT_CN):
    run.font.name = name
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(tag, rpr.nsmap)
        if el is None:
            el = OxmlElement(tag)
            rpr.append(el)
        el.set("typeface", name)


def clear_frame(tf):
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = Inches(0)
    tf.margin_right = Inches(0)
    tf.margin_top = Inches(0)
    tf.margin_bottom = Inches(0)


def strip_placeholder(shape):
    try:
        # Freeze inherited placeholder geometry before removing the placeholder marker.
        left, top, width, height = shape.left, shape.top, shape.width, shape.height
        shape.left = left
        shape.top = top
        shape.width = width
        shape.height = height
        nvpr = shape._element.nvSpPr.nvPr
        ph = nvpr.find(qn("p:ph"))
        if ph is not None:
            nvpr.remove(ph)
    except Exception:
        pass


def add_para(tf, text="", first=False, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT,
             level=0, space_after=6, space_before=0, line_spacing=1.5, font_name=FONT_CN):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.level = level
    p.alignment = align
    p.space_after = Pt(space_after)
    p.space_before = Pt(space_before)
    if line_spacing is not None:
        p.line_spacing = line_spacing
    r = p.add_run()
    r.text = text
    set_run_font(r, size=size, bold=bold, name=font_name)
    r.font.color.rgb = color
    return p


def fill_title(slide, text, size=30):
    for sh in slide.shapes:
        if getattr(sh, "is_placeholder", False):
            try:
                if sh.placeholder_format.type == 1:  # TITLE
                    tf = sh.text_frame
                    clear_frame(tf)
                    add_para(tf, text, first=True, size=size, bold=True, color=INK,
                             line_spacing=None, space_after=0)
                    strip_placeholder(sh)
                    return sh
            except Exception:
                pass
    box = slide.shapes.add_textbox(Inches(0.67), Inches(0.55), Inches(12.0), Inches(0.8))
    tf = box.text_frame
    clear_frame(tf)
    add_para(tf, text, first=True, size=size, bold=True, color=INK,
             line_spacing=None, space_after=0)
    return box


def body_placeholder(slide):
    for sh in slide.shapes:
        if getattr(sh, "is_placeholder", False):
            try:
                if sh.placeholder_format.type in (2, 7):  # BODY or OBJECT
                    return sh
            except Exception:
                pass
    return None


def fill_body(slide, lines, size=16.5, color=INK, space_after=8):
    sh = body_placeholder(slide)
    if sh is None:
        return None
    tf = sh.text_frame
    clear_frame(tf)
    first = True
    for item in lines:
        if isinstance(item, str):
            text, opts = item, {}
        else:
            text, opts = item
        add_para(
            tf,
            text,
            first=first,
            size=opts.get("size", size),
            bold=opts.get("bold", False),
            color=opts.get("color", color),
            level=opts.get("level", 0),
            space_after=opts.get("space_after", space_after),
            space_before=opts.get("space_before", 0),
            line_spacing=opts.get("line_spacing", 1.5),
        )
        first = False
    strip_placeholder(sh)
    return sh


def fill_comparison(slide, left_header, left_lines, right_header, right_lines,
                    left_size=16, right_size=16):
    body_shapes = []
    header_shapes = []
    for sh in slide.shapes:
        if not getattr(sh, "is_placeholder", False):
            continue
        try:
            typ = sh.placeholder_format.type
            idx = sh.placeholder_format.idx
            if typ == 2:
                header_shapes.append((idx, sh))
            elif typ == 7:
                body_shapes.append((idx, sh))
        except Exception:
            pass
    if len(header_shapes) < 2 or len(body_shapes) < 2:
        return
    header_shapes.sort(key=lambda x: x[0])
    body_shapes.sort(key=lambda x: x[0])
    for sh, text in ((header_shapes[0][1], left_header), (header_shapes[1][1], right_header)):
        tf = sh.text_frame
        clear_frame(tf)
        add_para(tf, text, first=True, size=17, bold=True, color=ACCENT,
                 line_spacing=None, space_after=0)
        strip_placeholder(sh)
    for sh, lines, size in (
        (body_shapes[0][1], left_lines, left_size),
        (body_shapes[1][1], right_lines, right_size),
    ):
        tf = sh.text_frame
        clear_frame(tf)
        first = True
        for item in lines:
            if isinstance(item, str):
                text, opts = item, {}
            else:
                text, opts = item
            add_para(
                tf,
                text,
                first=first,
                size=opts.get("size", size),
                bold=opts.get("bold", False),
                color=opts.get("color", INK),
                space_after=opts.get("space_after", 7),
                line_spacing=1.5,
            )
            first = False
        strip_placeholder(sh)


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def add_page_number(slide, n, total):
    box = slide.shapes.add_textbox(Inches(12.05), Inches(7.12), Inches(0.62), Inches(0.25))
    tf = box.text_frame
    clear_frame(tf)
    add_para(tf, f"{n:02d}", first=True, size=12, bold=True, color=SUB,
             align=PP_ALIGN.RIGHT, line_spacing=None, space_after=0, font_name=FONT_EN)


def add_text(slide, text, x, y, w, h, size=16, bold=False, color=INK,
             align=PP_ALIGN.LEFT, font=FONT_CN, line_spacing=1.5, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    clear_frame(tf)
    tf.vertical_anchor = valign
    add_para(tf, text, first=True, size=size, bold=bold, color=color,
             align=align, line_spacing=line_spacing, space_after=0, font_name=font)
    return box


def add_metric(slide, number, label, x, y, w,
               number_size=30, label_size=14, color=ACCENT):
    add_text(slide, number, x, y, w, 0.62, size=number_size, bold=True, color=color,
             font=FONT_EN, line_spacing=None, valign=MSO_ANCHOR.BOTTOM)
    add_text(slide, label, x, y + 0.69, w, 0.58, size=label_size, color=SUB,
             line_spacing=1.5)


def add_card(slide, x, y, w, h, fill=PALE, line=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(0.75)
    return shp


def add_citation(slide, text, y=6.94):
    add_text(slide, text, 10.15, y, 2.5, 0.26, size=12, color=SUB,
             align=PP_ALIGN.RIGHT, line_spacing=None)


def set_section_slide(slide, number, title, lead, points):
    # Layout 2 contains title and body placeholders near the vertical center.
    title_sh = None
    body_sh = None
    for sh in slide.shapes:
        if not getattr(sh, "is_placeholder", False):
            continue
        try:
            if sh.placeholder_format.type == 1:
                title_sh = sh
            elif sh.placeholder_format.type == 2:
                body_sh = sh
        except Exception:
            pass
    if title_sh is not None:
        tf = title_sh.text_frame
        clear_frame(tf)
        add_para(tf, title, first=True, size=30, bold=True, color=INK,
                 line_spacing=None, space_after=0)
        strip_placeholder(title_sh)
    if body_sh is not None:
        body_sh._element.getparent().remove(body_sh._element)

    add_text(slide, lead, 2.18, 5.05, 8.5, 0.58,
             size=16, color=SUB, line_spacing=1.5)
    for i, point in enumerate(points):
        add_text(slide, f"0{i + 1}", 2.18, 5.78 + i * 0.43, 0.5, 0.34,
                 size=13, bold=True, color=ACCENT, font=FONT_EN, line_spacing=None)
        add_text(slide, point, 2.78, 5.76 + i * 0.43, 7.9, 0.36,
                 size=14, color=INK, line_spacing=1.5)

    add_text(slide, number, 0.8, 1.15, 1.0, 0.5, size=24, bold=True,
             color=ACCENT, font=FONT_EN, line_spacing=None)


def add_footer_rule(slide):
    # The source template has no decorative rules; retain that visual language.
    return None


# Patch unavailable template fonts before opening the file.
patch_package_fonts(BASE, PATCHED)
prs = Presentation(PATCHED)

# Existing template pages: title, first-act divider, Card method, Card result,
# Card results continuation, Fogg/Wairagkar comparison.
slides = list(prs.slides)

# ---------------------------------------------------------------------------
# 01 · Title
# ---------------------------------------------------------------------------
s = slides[0]
fill_title(s, "从解码到可用：脑机接口功能重建的证据链", size=32)
# subtitle placeholder
for sh in s.shapes:
    if getattr(sh, "is_placeholder", False):
        try:
            if sh.placeholder_format.type == 4:
                tf = sh.text_frame
                clear_frame(tf)
                add_para(tf, "2024–2026 年语言、运动、视觉、听觉与意识研究", first=True,
                         size=20, color=SUB, line_spacing=1.5, space_after=8)
                add_para(tf, "文献汇报｜38 篇精读材料，20 篇核心证据", size=14,
                         color=SUB, line_spacing=1.5, space_after=0)
                strip_placeholder(sh)
                sh.height = Inches(1.05)
        except Exception:
            pass

gate_y = 5.12
gate_titles = [
    ("01  可解码", "语音、运动、光幻视与隐蔽认知"),
    ("02  可独立使用", "居家、少标定、跨月稳定"),
    ("03  临床获益", "对照、长期安全、临床终点"),
]
for i, (h, body) in enumerate(gate_titles):
    x = 1.22 + i * 3.72
    add_card(s, x, gate_y, 3.28, 1.18, fill=PALE)
    add_text(s, h, x + 0.2, gate_y + 0.15, 2.85, 0.35, size=16, bold=True, color=ACCENT,
             line_spacing=None)
    add_text(s, body, x + 0.2, gate_y + 0.57, 2.85, 0.38, size=13, color=SUB, line_spacing=1.5)

add_text(s, "核心判断：技术竞争已经从单点解码指标转向长期稳定性、个体化成本和临床证据。",
         1.22, 6.55, 10.9, 0.40, size=14, bold=True, color=INK, line_spacing=1.5)
add_notes(
    s,
    "开场先说明：这不是五方向资料汇编，而是一条证据链。过去两年语音、运动、视觉和意识检测都出现了重要进展，但它们回答的问题不同。"
    "本汇报按“可解码—可独立使用—临床获益可归因”三层展开，并在每个方向保留单例、单臂、回顾性研究等限制。"
)

# ---------------------------------------------------------------------------
# 02 · Act 1 divider
# ---------------------------------------------------------------------------
s = slides[1]
set_section_slide(
    s,
    "01",
    "从“能解”到“能用”",
    "BCI 是否能从实验室信号解码，进入长期、低负担和真实的患者生活？",
    [
        "长期使用：Card 2026，超过 3,800 小时居家使用",
        "高通量：Jude 2026，最高 110 字符/分",
        "持久修复：Chandrasekaran 2026，关机后增益超过 2 个月",
    ],
)
add_notes(
    s,
    "第一幕关注语言与运动。这里的转折不是“又完成了一个解码任务”，而是研究问题开始变成：患者能否不依赖研究者、能否跨日稳定、能否在系统关闭后仍留下功能改善。"
    "三个入口数字分别代表可用性、交互通量和持久修复。"
)

# ---------------------------------------------------------------------------
# 03 · Act 1 framing
# ---------------------------------------------------------------------------
s = slides[2]
fill_title(s, "第一幕的三个问题：长时使用、规模迁移与持久修复")
body_main = fill_body(s, [
    ("1. 高通量能否进入真实生活？", {"size": 20, "bold": True, "color": ACCENT}),
    "Card 2026 把评价指标从实验室词错误率转向无人协助的居家电脑操作。",
    ("2. 模型能否从一个患者迁移到另一个患者？", {"size": 20, "bold": True, "color": ACCENT}),
    "Fogg 2026 用 6 人联合训练，探索新患者少于 200 句微调即可使用的路线。",
    ("3. 系统关闭后是否仍留下治疗效应？", {"size": 20, "bold": True, "color": ACCENT}),
    "Chandrasekaran 2026 把“替代”与“修复”拆成两个不同终点。",
], size=17)
add_citation(s, "[1][2][6]")
add_notes(
    s,
    "这一页先给第一幕搭问题框架。第一节不是简单介绍 Card，而是说明语音 BCI 的评价对象从模型性能转向患者生活；第二节讨论 Fogg 的跨用户迁移；第三节讨论 Chandrasekaran 提出的关机后残留增益。"
)

# ---------------------------------------------------------------------------
# 04 · Card 2026 — question, design, method
# ---------------------------------------------------------------------------
s = slides[3]
fill_title(s, "Card 2026｜问题与设计：把语音 BCI 搬回家")
body_main = fill_body(s, [
    ("Q：皮层内语音 BCI 能否在无人协助的家庭环境中长期运行？", {"size": 20, "bold": True, "color": ACCENT}),
    "受试者为一名 ALS 相关的重度构音障碍、四肢瘫痪患者。",
    "植入皮层内微电极阵列，覆盖语言相关腹侧前中央回。",
    "语音通道：transformer 脑到文本解码器，把尝试说话转成键盘输入。",
    "光标通道：RNN 神经光标解码器，把运动意图转成鼠标移动。",
    "BG Home 中间层把两条通道映射为个人电脑的键鼠事件。",
], size=16.5)
body_main.width = Inches(7.65)

# small flow on the right
flow_x = 8.72
flow_y = 2.22
add_text(s, "一条植入，两条输出", flow_x, 1.78, 3.6, 0.38,
         size=15, bold=True, color=ACCENT)
for i, (label, sub) in enumerate([
    ("尝试说话", "vPCG 神经活动"),
    ("文本 + 光标", "两个解码器"),
    ("个人电脑", "邮件、网页、通话"),
]):
    yy = flow_y + i * 1.05
    add_card(s, flow_x, yy, 3.55, 0.78, fill=PALE_BLUE if i == 0 else PALE)
    add_text(s, label, flow_x + 0.18, yy + 0.10, 1.3, 0.38, size=14, bold=True, color=INK)
    add_text(s, sub, flow_x + 1.45, yy + 0.10, 1.9, 0.40, size=12, color=SUB)
    if i < 2:
        add_text(s, "↓", flow_x + 1.55, yy + 0.80, 0.45, 0.32, size=16,
                 color=ACCENT, align=PP_ALIGN.CENTER, line_spacing=None)

add_citation(s, "[1]")
add_notes(
    s,
    "Card 2026 的研究问题不是再次展示语音解码，而是验证无人协助的居家使用。受试者为 ALS 相关重度构音障碍患者，系统同时输出文本和光标。"
    "语音通道使用 transformer，光标通道使用 RNN，最后由中间层映射为普通电脑的键盘和鼠标事件。这个设计决定了后文的 3,800 小时和 56 词/分不是实验室内短时指标。"
)

# ---------------------------------------------------------------------------
# 05 · Card 2026 — results and limits
# ---------------------------------------------------------------------------
s = slides[4]
fill_title(s, "Card 2026｜结果与边界：真实生活证据，但仍是单例")

# Clear original body placeholder and use stat blocks.
body_sh = body_placeholder(s)
if body_sh is not None:
    body_sh._element.getparent().remove(body_sh._element)

metrics = [
    ("3,800 h", "近 2 年累计居家使用，444/653 天出现使用"),
    ("56 wpm", "自由使用平均速度；18.3 万句、196 万词"),
    ("92%", "患者自评句子至少基本正确"),
]
for i, (num, lab) in enumerate(metrics):
    x = 0.72 + i * 4.02
    add_metric(s, num, lab, x, 1.55, 3.55, number_size=34, label_size=13)

add_card(s, 0.7, 3.25, 5.65, 1.7, fill=PALE)
add_text(s, "超过 99% 的任务准确率", 0.93, 3.42, 4.95, 0.45, size=18, bold=True, color=INK)
add_text(s, "来自屏幕提示的复制任务，词表 125,000；不能等同于自由对话准确率。",
         0.93, 3.94, 4.95, 0.72, size=14, color=SUB)

add_card(s, 6.65, 3.25, 5.75, 1.7, fill=RGBColor(0xFA, 0xEF, 0xEF))
add_text(s, "必须同时说的限制", 6.88, 3.42, 5.05, 0.45, size=18, bold=True, color=ALERT)
add_text(s, "单例 ALS；自评无独立金标准；原始神经数据与居家文本因隐私不公开。",
         6.88, 3.94, 5.05, 0.72, size=14, color=SUB)

add_text(s,
         "结论：语音 BCI 的主指标已经从 WER 转向“患者能否无人协助地生活”，但单例成功还不能代表可推广的临床系统。",
         0.72, 5.58, 11.75, 0.66, size=16, bold=True, color=INK)
add_citation(s, "[1]")
add_notes(
    s,
    "卡片的三个核心结果是 3,800 小时、56 词/分和 92% 自评。系统在家输出约 18.3 万句、196 万词，患者还能发送邮件、浏览网页、视频通话并维持工作。"
    "但页面必须同时保留两个限制：超过 99% 是提示词表任务，92% 是患者自评；研究只有一名患者，原始神经数据与居家转录也无法公开。"
)

# ---------------------------------------------------------------------------
# 06 · Fogg + Wairagkar
# ---------------------------------------------------------------------------
s = slides[5]
fill_title(s, "Fogg + Wairagkar｜规模化与表达力")
fill_comparison(
    s,
    "Fogg 2026（预印本）：一个模型服务多个用户",
    [
        "6 名皮层内语音 BCI 参与者，覆盖不同性别、病因与说话策略。",
        "多用户 transformer 对每个人都优于单用户模型。",
        "相对 WER 平均下降超过 50%。",
        "完全留出的新用户少于 200 句微调后，WER 低于 7%。",
        "边界：同类 NeuroPort 阵列，语料以会话英语为主。",
    ],
    "Wairagkar 2025：从文本走向有语气的声音",
    [
        "256 个皮层内电极，覆盖腹侧前中央回。",
        "直接把神经活动映射为声学特征，而非先转文字。",
        "问句语调 90.5%，词重音 95.7%。",
        "闭环音频反馈；可实时改变语调并演唱短旋律。",
        "边界：单例；精确控制音高对患者有额外认知负担。",
    ],
    left_size=15,
    right_size=15,
)
add_text(s, "共同意义：语音 BCI 的竞争维度从单一 WER 扩展为数据效率、跨用户泛化与表达能力。",
         0.78, 6.55, 11.65, 0.38, size=15, bold=True, color=INK)
add_citation(s, "[2][3]")
add_notes(
    s,
    "Fogg 解决问题规模：过去每个患者要从数千句训练数据开始，现在 6 人联合训练后，新用户少于 200 句即可达到低于 7% WER。"
    "Wairagkar 解决表达力：256 个电极直接预测声学特征，保留语调、重音和音色。Fogg 的限制是同类硬件和英语语料，Wairagkar 是单例。"
)

# Add new slides from this point.


def new_content_slide(title):
    sl = prs.slides.add_slide(prs.slide_layouts[1])
    fill_title(sl, title)
    return sl


def new_compare_slide(title):
    sl = prs.slides.add_slide(prs.slide_layouts[4])
    fill_title(sl, title)
    return sl


def new_section_slide(number, title, lead, points):
    sl = prs.slides.add_slide(prs.slide_layouts[2])
    set_section_slide(sl, number, title, lead, points)
    return sl


# ---------------------------------------------------------------------------
# 07 · Jude + Karpowicz
# ---------------------------------------------------------------------------
s = new_compare_slide("Jude + Karpowicz｜运动 BCI 的通量与稳定性")
fill_comparison(
    s,
    "Jude 2026：把双手 QWERTY 做成神经假体",
    [
        "2 名参与者；尝试双手手指屈伸，皮层内阵列解码后输出 QWERTY 字符。",
        "最高 110 字符/分（约 22 词/分），词错误率 1.6%。",
        "30 句标定即可达到主要性能；blank 类让用户自行决定输出节奏。",
        "日间漂移真实存在；逐日在线训练可降低错误率。",
        "边界：两例，最优参与者指标，提示性句子为主。",
    ],
    "Karpowicz 2025：NoMAD 让解码器跨月稳定",
    [
        "核心问题是神经信号非平稳，昨日的解码器今天可能失效。",
        "NoMAD 用 LFADS 推断潜在动力学，再做无监督流形对齐。",
        "猴运动皮层上稳定至少 3 个月，不需要有监督重标定。",
        "ADAN 半衰期仅 7.12 天并出现 15 次失败；Aligned FA 约 45.1 天。",
        "边界：长期结果主要来自猴；人类为闭环光标演示。",
    ],
    left_size=15,
    right_size=15,
)
add_text(s, "一对关系：Jude 证明“能高速输入”，Karpowicz 解决“能否每天都稳定使用”。",
         0.78, 6.55, 11.65, 0.38, size=15, bold=True, color=INK)
add_citation(s, "[4][5]")
add_notes(
    s,
    "Jude 2026 的关键是范式：利用患者熟悉的 QWERTY 与肌肉记忆，110 字符/分、1.6% WER、30 句标定。"
    "Karpowicz 2025 的核心则是稳定性：NoMAD 不重新学习脑—行为映射，而是把漂移后的神经活动对齐回稳定潜在空间。长期证据主要是猴，人类只做演示。"
)

# ---------------------------------------------------------------------------
# 08 · Chandrasekaran + Darvishi
# ---------------------------------------------------------------------------
s = new_compare_slide("Chandrasekaran + Darvishi｜从替代到修复")
fill_comparison(
    s,
    "Chandrasekaran 2026：双重神经旁路",
    [
        "1 名 C4 感觉/C5 运动 ASIA A 完全性脊髓损伤患者。",
        "M1 BCI 解码运动意图，脊髓刺激激活手部肌肉。",
        "S1 微刺激提供触觉反馈，与运动意图同步形成“cortical mirroring”。",
        "恢复自主进食与精细操作；关机后肘屈曲和腕部触觉改善 >2 个月。",
        "边界：单例、开放标签；BCI/SCS/皮层刺激无法拆分归因。",
    ],
    "Darvishi 2024：RehabSwift 非侵入康复",
    [
        "12 名慢性期卒中幸存者，18 次个体化运动想象 BCI 训练。",
        "FMA-UE：36.75 → 45.08；随访 46.17。",
        "12 人中 9 人达到 MCID，效果至少维持 4 周。",
        "客观改善明显，但主观目标达成率只有 60%。",
        "边界：单臂、无对照、随访仅 4–6 周，不能单独归因于 BCI。",
    ],
    left_size=15,
    right_size=15,
)
add_card(s, 0.78, 6.28, 11.72, 0.72, fill=PALE)
add_text(s, "本汇报推论：神经假体的终点应从“开机时能做什么”扩展到“关机后还剩下什么”。",
         1.0, 6.42, 11.25, 0.42, size=15, bold=True, color=ACCENT)
add_citation(s, "[6][7]")
add_notes(
    s,
    "Chandrasekaran 把 BCI、脊髓刺激和皮层感觉反馈组合在同一系统。最关键的是关机后仍保留肘屈曲与腕部触觉改善，但三种干预无法拆分，且只是单例。"
    "Darvishi 是另一极：非侵入、可规模化，12 人 FMA-UE 明显改善，但无对照、随访短。合起来提出新的证据问题：替代效应与修复效应如何分开？"
)

# ---------------------------------------------------------------------------
# 09 · Act 2 divider
# ---------------------------------------------------------------------------
s = new_section_slide(
    "02",
    "从“替代器件”到“闭环控制”",
    "视觉已经出现临床终点，听觉器件最成熟，但两者的闭环控制都还没有完成。",
    [
        "视觉临床终点：PRIMA 81.3% 达到视力改善",
        "视觉双向预测：Grani 2025 在 V1 刺激并记录",
        "听觉闭环边界：Haro 2025 只显著抑制干扰者",
    ],
)
add_notes(
    s,
    "第二幕展示成熟度错位。视觉已经有像临床器械试验一样的终点，但主要来自视网膜路线；皮层路线仍处双向预测阶段。听觉的人工耳蜗和预后预测很成熟，但真正闭环听觉注意控制仍只有小样本、短时实验。"
)

# ---------------------------------------------------------------------------
# 10 · PRIMA + Petoe
# ---------------------------------------------------------------------------
s = new_compare_slide("视觉重建｜临床终点与长期功能结局")
fill_comparison(
    s,
    "PRIMA 2026：第一次出现临床主要终点",
    [
        "38 例 AMD 地图样萎缩；视网膜下 2×2 mm 光伏阵列，100 µm 像素。",
        "12 个月时 81.3% 达到 ≥logMAR 0.2 改善。",
        "平均改善 24.5–25.5 个字母；logMAR 1.32 对应像素采样极限。",
        "单臂、开放标签、基线对照；6 人未完成 12 个月评估。",
        "意义：人工视觉有了可跨路线比较的临床基准。",
    ],
    "Petoe 2025：长期戴着假体生活是什么样",
    [
        "4 例终末期视网膜色素变性；44 通道脉络膜上假体。",
        "随访 2.7 年，无装置相关严重不良事件。",
        "4/4 定位改善；3/4 运动辨别；2/4 空间与物体识别改善。",
        "出现无人协助探索新环境、发现他人和避障。",
        "边界：4 例、单臂、非遮蔽，患者能识别关机条件。",
    ],
    left_size=15,
    right_size=15,
)
add_text(s, "两页功课：PRIMA 回答“视力是否改善”，Petoe 回答“长期生活是否受益”，两者都尚未解决实时闭环。",
         0.78, 6.55, 11.65, 0.4, size=15, bold=True, color=INK)
add_citation(s, "[8][9]")
add_notes(
    s,
    "PRIMA 是视觉部分的核心基准：38 例、81.3%、24.5–25.5 个字母，并与 100 µm 像素尺寸对应。"
    "Petoe 提供 2.7 年功能与安全随访，说明长期使用有价值；但只有四例，且非遮蔽。两篇都说明视觉假体已经进入临床评价，但还没有实时读取感知反馈。"
)

# ---------------------------------------------------------------------------
# 11 · Grani
# ---------------------------------------------------------------------------
s = new_content_slide("Grani 2025｜皮层视觉：从刺激走向双向预测")
body_sh = body_placeholder(s)
if body_sh is not None:
    body_sh._element.getparent().remove(body_sh._element)

add_text(s, "刺激参数", 0.82, 1.58, 2.5, 0.46, size=18, bold=True, color=ACCENT,
         align=PP_ALIGN.CENTER)
add_text(s, "邻近神经活动", 4.35, 1.58, 3.4, 0.46, size=18, bold=True, color=ACCENT,
         align=PP_ALIGN.CENTER)
add_text(s, "主观光幻视", 8.55, 1.58, 3.1, 0.46, size=18, bold=True, color=ACCENT,
         align=PP_ALIGN.CENTER)
add_text(s, "→", 3.28, 1.62, 1.05, 0.4, size=22, color=SUB, align=PP_ALIGN.CENTER,
         line_spacing=None)
add_text(s, "→", 7.7, 1.62, 0.85, 0.4, size=22, color=SUB, align=PP_ALIGN.CENTER,
         line_spacing=None)

add_card(s, 0.82, 2.25, 5.72, 2.42, fill=PALE)
add_text(s, "研究设计", 1.05, 2.40, 5.12, 0.44, size=17, bold=True, color=INK)
for i, line in enumerate([
    "2 名盲人志愿者；右侧枕极 V1 植入 100 电极。",
    "刺激邻近电极同步记录神经活动。",
    "假刺激试次用于估计误报率。",
    "阵列 6 个月后取出。",
]):
    add_text(s, line, 1.05, 2.93 + i * 0.43, 5.12, 0.36, size=14, color=INK)

add_card(s, 6.78, 2.25, 5.75, 2.42, fill=PALE_BLUE)
add_text(s, "核心结果", 7.02, 2.40, 5.15, 0.44, size=17, bold=True, color=INK)
for i, line in enumerate([
    "阈值、亮度和光幻视数量可由神经活动预测。",
    "两次独立光幻视的最小间隔约 0.30 / 0.33 秒。",
    "低误报率支持感知由刺激引起。",
    "预测是离线统计推断，尚非实时闭环控制。",
]):
    add_text(s, line, 7.02, 2.93 + i * 0.43, 5.15, 0.36, size=14, color=INK)

add_text(s,
         "下一步：用人体“刺激—神经活动—主观感知”配对数据校准图像到光幻视的正向模型。",
         0.82, 5.1, 11.7, 0.58, size=16, bold=True, color=INK)
add_text(s,
         "本汇报推论：皮层视觉的关键瓶颈正在从电极数量转向可控性与模型校准。",
         0.82, 5.82, 11.7, 0.42, size=14, color=SUB)
add_citation(s, "[10]")
add_notes(
    s,
    "Grani 2025 纳入两名盲人，在 V1 植入 100 电极，同时刺激和记录。研究证明阈值、亮度和光幻视数量可由神经活动预测，并给出两次光幻视区分的时间阈值。"
    "但这是两例、六个月、离线预测，不是实时闭环控制。它的真正价值是提供人体配对数据，可用于校准算法层的正向模型。"
)

# ---------------------------------------------------------------------------
# 12 · Hearing
# ---------------------------------------------------------------------------
s = new_content_slide("听觉重建｜临床预测、AI 落地与闭环控制是三种成熟度")
body_sh = body_placeholder(s)
if body_sh is not None:
    body_sh._element.getparent().remove(body_sh._element)

lanes = [
    (0.72, "临床预测", "Wang 2026", [
        "278 名儿童，3 国、3 语言",
        "术前 T1 MRI 预测语言结局",
        "准确率 92.39%，AUC 0.98",
        "纯临床模型仅 53.57%",
        "边界：内部验证为主，未进入临床流程",
    ]),
    (4.46, "AI 落地", "Nair 2025", [
        "系统综述纳入 20 项研究",
        "影像预测与语音增强有前景",
        "没有任何模型进入 CI 编程流程",
        "缺口是工程、监管与流程整合",
    ]),
    (8.2, "真正闭环", "Haro 2025", [
        "22 名听力正常者，单次 1 小时",
        "训练阶段平均解码 60.8%",
        "干扰者神经追踪显著下降",
        "目标者追踪未显著增强",
        "健康人、湿电极、非自适应解码器",
    ]),
]
for x, head, paper, lines in lanes:
    add_card(s, x, 1.62, 3.38, 4.75, fill=PALE)
    add_text(s, head, x + 0.2, 1.80, 2.95, 0.45, size=18, bold=True, color=ACCENT)
    add_text(s, paper, x + 0.2, 2.28, 2.95, 0.36, size=13, color=SUB)
    for i, line in enumerate(lines):
        add_text(s, line, x + 0.2, 2.72 + i * 0.63, 2.95, 0.70,
                 size=13.5, color=INK, line_spacing=1.5)
add_text(s, "总判断：成熟器件 ≠ 成熟闭环；闭环听觉的作用点仍可能是“抑制干扰”，而非“增强目标”。",
         0.78, 6.5, 11.7, 0.38, size=15, bold=True, color=INK)
add_citation(s, "[11][12][13]")
add_notes(
    s,
    "听觉要区分三种成熟度。Wang 2026 用 278 名儿童证明术前 MRI 有很强预测力；Nair 2025 发现 20 项 AI 研究中没有模型真正进入人工耳蜗编程；Haro 2025 的无创闭环只显著抑制干扰者，目标追踪未增强。"
    "所以人工耳蜗器件成熟，不等于闭环听觉 BCI 成熟。"
)

# ---------------------------------------------------------------------------
# 13 · Act 3 divider
# ---------------------------------------------------------------------------
s = new_section_slide(
    "03",
    "从“有没有意识”到“能否证明内容”",
    "可检测、可干预、可解码和可证明是四个不同层级。",
    [
        "检测：241 名无指令反应者中 25% 检出 CMD",
        "促醒：DBS 应答约 44.7%，但仍无 RCT",
        "内容：语义可重建，但主观体验没有 ground truth",
    ],
)
add_notes(
    s,
    "第三幕按四层推进：状态检测、指令执行、语义内容、主观体验。前两层已有较明确的临床与 BCI 证据；第三层在健康人中有原理验证；第四层仍没有可靠的行为金标准。"
)

# ---------------------------------------------------------------------------
# 14 · Bodien + du Bois
# ---------------------------------------------------------------------------
s = new_compare_slide("意识检测｜行为无反应不等于没有认知")
fill_comparison(
    s,
    "Bodien 2024：发现认知—运动分离（CMD）",
    [
        "353 名意识障碍患者，6 个国际中心，2006–2023。",
        "241 名行为上无指令反应者接受任务态 EEG / fMRI。",
        "60 名检出 CMD，占 25%；11 名 fMRI、13 名 EEG、36 名联合阳性。",
        "CMD 与更年轻、病程更长、创伤性病因相关。",
        "边界：单次阴性不能证明无意识；任务理解与疲劳可导致假阴性。",
    ],
    "du Bois 2026：BCI 成为互补诊断工具",
    [
        "44 名参与者：UWS 14、MCS 17、LIS 11、健康对照 2。",
        "可穿戴 EEG 完成运动想象、反馈训练与二选一问答。",
        "73.8% 获得显著解码；90% 进入问答测试。",
        "融合 BCI 与行为评估后，MCS 灵敏度 39%→69%。",
        "边界：BCI 阴性可能来自电极、伪迹、疲劳或任务失败。",
    ],
    left_size=15,
    right_size=15,
)
add_text(s, "BCI 的定位不是替代临床诊断，而是成为可重复、无需运动输出的认知探针。",
         0.78, 6.55, 11.65, 0.38, size=15, bold=True, color=INK)
add_citation(s, "[14][15]")
add_notes(
    s,
    "Bodien 2024 是临床问题起点：353 名患者中，241 名没有可观察的指令反应，但 25% 检出 CMD。"
    "du Bois 2026 把这一思路做成可穿戴 EEG 多阶段 BCI：73.8% 有显著解码，BCI 与行为评估融合后 MCS 灵敏度从 39% 到 69%。关键限制是 BCI 失败不能直接解释为无意识。"
)

# ---------------------------------------------------------------------------
# 15 · DoC neuromodulation
# ---------------------------------------------------------------------------
s = new_content_slide("意识促醒｜靶点、机制与证据等级")
body_sh = body_placeholder(s)
if body_sh is not None:
    body_sh._element.getparent().remove(body_sh._element)

cols = [
    (0.72, "哪里刺激", "Warren 2025 · n=40", [
        "回顾性 DBS 队列 + 电场建模",
        "有效电场越过 CM，覆盖 PF 下核与 VTT",
        "纹状体灰质保留更好者更可能改善",
        "关键体素未通过 FDR 校正",
    ]),
    (4.46, "为什么可能有效", "Zhang 2025 · n=23", [
        "中央丘脑直接电生理记录",
        "θ 节律跨病因解释恢复差异",
        "识别“早期差、最终恢复”亚组",
        "已决定植入者存在选择偏倚",
    ]),
    (8.2, "证据到底够不够", "Zhou 2026 · 综述", [
        "ARAS—中环路—ICN 三层框架",
        "CM-Pf/CL DBS 应答约 44.7%",
        "仍无 RCT；唯一 Level 1 来自药物",
        "TIS 仅有 9 例、10 天试点",
    ]),
]
for x, head, paper, lines in cols:
    add_card(s, x, 1.66, 3.38, 4.05, fill=PALE)
    add_text(s, head, x + 0.2, 1.82, 2.98, 0.45, size=18, bold=True, color=ACCENT)
    add_text(s, paper, x + 0.2, 2.30, 2.98, 0.36, size=13, color=SUB)
    for i, line in enumerate(lines):
        add_text(s, line, x + 0.2, 2.82 + i * 0.72, 2.98, 0.70,
                 size=14, color=INK)
add_text(s, "下一步不是继续累积单臂病例，而是按病因、病程、CRS-R 与网络完整性做分层随机试验。",
         0.78, 5.95, 11.7, 0.38, size=15, bold=True, color=INK)
add_text(s, "本汇报推论：促醒疗效需要与刺激层级匹配的终点，而不是所有患者共享一个平均分数。",
         0.78, 6.44, 11.7, 0.34, size=13, color=SUB)
add_citation(s, "[16][17][18]")
add_notes(
    s,
    "Warren 通过 40 例回顾性队列说明有效电场需要覆盖 PF 下核和 VTT；Zhang 用 23 例中央丘脑直接记录发现 θ 节律与恢复相关；Zhou 的综述则给出整体证据坐标：DBS 应答约 44.7%，但无 RCT，唯一 Level 1 证据来自药物。"
    "所以这一领域有机制和靶点进展，但疗效尚未被高质量随机试验确立。"
)

# ---------------------------------------------------------------------------
# 16 · Content decoding
# ---------------------------------------------------------------------------
s = new_content_slide("意识内容｜可解码、可重建语义，不等于可证明体验")
body_sh = body_placeholder(s)
if body_sh is not None:
    body_sh._element.getparent().remove(body_sh._element)

levels = [
    ("01", "状态检测", "UWS / MCS / LIS / CMD", "Bodien；du Bois"),
    ("02", "指令执行", "运动想象、yes/no、命令跟随", "du Bois"),
    ("03", "语义内容", "词语、句子与概念重建", "Tang"),
    ("04", "主观体验", "患者此刻真正体验到了什么", "尚无金标准"),
]
for i, (num, head, body, evidence) in enumerate(levels):
    y = 1.64 + i * 0.9
    fill = PALE_BLUE if i < 2 else (PALE if i == 2 else RGBColor(0xFA, 0xEF, 0xEF))
    add_card(s, 0.72, y, 6.2, 0.72, fill=fill)
    add_text(s, num, 0.92, y + 0.18, 0.45, 0.32, size=13, bold=True,
             color=ACCENT if i < 3 else ALERT, font=FONT_EN, line_spacing=None)
    add_text(s, head, 1.5, y + 0.12, 1.4, 0.4, size=15, bold=True, color=INK,
             line_spacing=None)
    add_text(s, body, 2.95, y + 0.12, 2.4, 0.4, size=13, color=INK,
             line_spacing=None)
    add_text(s, evidence, 5.35, y + 0.12, 1.3, 0.4, size=12, color=SUB,
             align=PP_ALIGN.RIGHT, line_spacing=None)

add_card(s, 7.2, 1.64, 5.3, 1.52, fill=PALE)
add_text(s, "Tang 2023｜语义重建", 7.43, 1.76, 4.8, 0.44, size=16, bold=True, color=INK)
add_text(s, "fMRI 重建连续语言含义，而非逐字转录；可跨听觉、视觉与想象语言，但仍需个体化训练。",
         7.43, 2.23, 4.8, 0.72, size=13, color=SUB)

add_card(s, 7.2, 3.35, 5.3, 1.52, fill=PALE)
add_text(s, "Littlejohn 2025｜流式交流", 7.43, 3.47, 4.8, 0.44, size=16, bold=True, color=INK)
add_text(s, "80 ms 增量解码，约 46.5 词/分；仍需主动尝试说话。",
         7.43, 3.94, 4.8, 0.72, size=13, color=SUB)

add_card(s, 7.2, 5.00, 5.3, 1.15, fill=RGBColor(0xFA, 0xEF, 0xEF))
add_text(s, "Kronemer 2025｜证据学边界", 7.43, 5.10, 4.8, 0.45, size=15, bold=True, color=ALERT)
add_text(s, "必须区分意识水平、状态与内容；no-report、掩蔽与无意识对照不可省略。",
         7.43, 5.43, 4.8, 0.64, size=12.5, color=SUB)
add_citation(s, "[19][20][21]")
add_notes(
    s,
    "意识内容必须分四层讲。状态检测和指令执行已有临床 BCI 证据；Tang 证明 fMRI 可以重建语义，但不是逐字转录；Littlejohn 证明实时语言输出可达约 46.5 词/分。"
    "最后一层主观体验仍没有可靠 ground truth，Kronemer 的 no-report 与无意识对照是严格证明的核心。"
)

# ---------------------------------------------------------------------------
# 17 · Synthesis
# ---------------------------------------------------------------------------
s = new_content_slide("综合｜四条趋势与三个真正待解的问题")
body_sh = body_placeholder(s)
if body_sh is not None:
    body_sh._element.getparent().remove(body_sh._element)

trends = [
    ("01", "精度 → 可用性", "居家时长、标定成本、真实环境使用"),
    ("02", "个体模型 → 共享预训练", "跨用户语音、潜在动力学、电场建模"),
    ("03", "解剖核团 → 网络靶点", "丘脑—脑干—皮层通路与功能连接"),
    ("04", "开环刺激 → 闭环控制", "感觉反馈、声音反馈、注意反馈、再检测"),
]
for i, (num, head, body) in enumerate(trends):
    x = 0.78 + (i % 2) * 5.75
    y = 1.55 + (i // 2) * 1.42
    add_card(s, x, y, 5.32, 1.14, fill=PALE)
    add_text(s, num, x + 0.18, y + 0.18, 0.42, 0.28, size=12, bold=True,
             color=ACCENT, font=FONT_EN, line_spacing=None)
    add_text(s, head, x + 0.7, y + 0.12, 4.2, 0.45, size=16, bold=True, color=INK)
    add_text(s, body, x + 0.7, y + 0.58, 4.2, 0.34, size=13, color=SUB)

add_text(s, "留给组会的三个问题", 0.78, 4.49, 4.1, 0.45, size=17, bold=True, color=INK)
questions = [
    "关机后残留增益能否成为神经假体的独立终点？",
    "跨用户共享表征来自生物学共性，还是同一硬件阵列？",
    "意识内容何时能拥有不依赖行为的 ground truth？",
]
for i, q in enumerate(questions):
    add_text(s, f"{i+1}.", 0.82, 5.04 + i * 0.5, 0.35, 0.3, size=15, bold=True,
             color=ACCENT, font=FONT_EN, line_spacing=None)
    add_text(s, q, 1.25, 5.04 + i * 0.5, 6.1, 0.4, size=15, color=INK)

add_card(s, 8.0, 4.72, 4.52, 1.45, fill=PALE_BLUE)
add_text(s, "本汇报推论", 8.23, 4.92, 4.0, 0.32, size=13, bold=True, color=ACCENT)
add_text(s, "下一阶段竞争：长期稳定、个体化成本、闭环安全与临床终点同时成立。",
         8.23, 5.27, 4.0, 0.78, size=15, bold=True, color=INK)
add_citation(s, "[1][2][4][6][8][16][19]")
add_notes(
    s,
    "结论只保留四条趋势和三个问题。四条趋势分别是可用性、共享预训练、网络靶点和闭环控制。三个问题分别对应 Chandrasekaran 的关机后增益、Fogg 的共享表征来源，以及 Kronemer 的意识 ground truth。"
    "右侧是本汇报推论，不是任何单篇论文的结论。"
)


# ---------------------------------------------------------------------------
# 18–20 · Endnotes
# ---------------------------------------------------------------------------
references = [
    "[1] Card NS, Singer-Clark T, Peracha H, et al. Long-term independent use of an intracortical brain-computer interface for speech and cursor control. Nat Med. 2026;32(7):2504–2510. doi:10.1038/s41591-026-04414-6.",
    "[2] Fogg ZM, Card NS, Wairagkar M, et al. A generalizable speech neuroprosthesis. bioRxiv. 2026. doi:10.64898/2026.07.23.739430. [preprint]",
    "[3] Wairagkar M, Card NS, Singer-Clark T, et al. An instantaneous voice-synthesis neuroprosthesis. Nature. 2025;644(8075):145–152. doi:10.1038/s41586-025-09127-3.",
    "[4] Jude JJ, Levi-Aharoni H, Acosta AJ, et al. Restoring rapid natural bimanual typing with a neuroprosthesis after paralysis. Nat Neurosci. 2026. doi:10.1038/s41593-026-02218-y.",
    "[5] Karpowicz BM, Ali YH, Wimalasena LN, et al. Stabilizing brain-computer interfaces through alignment of latent dynamics. Nat Commun. 2025. doi:10.1038/s41467-025-59652-y.",
    "[6] Chandrasekaran S, Wandelt SK, Jangam A, et al. A neuroprosthesis for restoring hand movement and sensation in a person with complete tetraplegia. Nat Med. 2026. doi:10.1038/s41591-026-04498-0.",
    "[7] Darvishi S, Datta Gupta A, Hamilton-Bruce A, et al. Enhancing poststroke hand movement recovery: Efficacy of RehabSwift, a personalized brain-computer interface system. PNAS Nexus. 2024. doi:10.1093/pnasnexus/pgae240.",
    "[8] Holz FG, Le Mer Y, Muqit MMK, et al. Subretinal Photovoltaic Implant to Restore Vision in Geographic Atrophy Due to AMD. N Engl J Med. 2026. doi:10.1056/NEJMoa2501396.",
    "[9] Petoe MA, Abbott CJ, Titchener SA, et al. A Second-Generation (44-Channel) Suprachoroidal Retinal Prosthesis: A Single-Arm Clinical Trial of Feasibility. Ophthalmol Sci. 2025. doi:10.1016/j.xops.2024.100525.",
    "[10] Grani F, Soto-Sánchez C, Rodil Doblado A, et al. Neural correlates of phosphene perception in blind individuals: A step toward a bidirectional cortical visual prosthesis. Sci Adv. 2025. doi:10.1126/sciadv.adv8846.",
    "[11] Wang Y, Yuan D, Dettman S, et al. Forecasting Spoken Language Development in Children With Cochlear Implants Using Preimplant Magnetic Resonance Imaging. JAMA Otolaryngol Head Neck Surg. 2026. doi:10.1001/jamaoto.2025.4694.",
    "[12] Nair APS, Mishra SK, Andrade P, et al. A systematic review of machine learning approaches in cochlear implant outcomes. npj Digit Med. 2025;8:411. doi:10.1038/s41746-025-01733-9.",
    "[13] Haro S, Beauchene C, Quatieri TF, Smalt CJ. A Brain-Computer Interface for Improving Auditory Attention in Multi-Talker Environments. IEEE Access. 2025. doi:10.1109/ACCESS.2025.3623842.",
    "[14] Bodien YG, Allanson J, Cardone P, et al. Cognitive Motor Dissociation in Disorders of Consciousness. N Engl J Med. 2024;391(7):598–608. doi:10.1056/NEJMoa2400645.",
    "[15] du Bois N, Korik A, Hodge S, et al. Advancing EEG-based assessment of consciousness and cognition in prolonged disorders of consciousness. Commun Med. 2026;6:344. doi:10.1038/s43856-026-01574-x.",
    "[16] Warren AEL, Raguž M, Friedrich H, et al. A human brain network linked to restoration of consciousness after deep brain stimulation. Nat Commun. 2025. doi:10.1038/s41467-025-61988-4.",
    "[17] Zhang H, Ge Q, Liu X, et al. A shared central thalamus mechanism underlying diverse recoveries in disorders of consciousness. Nat Commun. 2025. doi:10.1038/s41467-025-65360-4.",
    "[18] Zhou Y, Cao T, Zhu S, et al. Neural Circuit Mechanisms of Disorders of Consciousness and Potential Therapeutic Targets of Brain Stimulation. CNS Neurosci Ther. 2026. doi:10.1002/cns.71037.",
    "[19] Kronemer SI, Bandettini PA, Gonzalez-Castillo J. Sleuthing subjectivity: a review of covert measures of consciousness. Nat Rev Neurosci. 2025;26(8):476–496. doi:10.1038/s41583-025-00934-1.",
    "[20] Tang J, LeBel A, Jain S, Huth AG. Semantic reconstruction of continuous language from non-invasive brain recordings. Nat Neurosci. 2023;26(5):858–866. doi:10.1038/s41593-023-01304-9.",
    "[21] Littlejohn KT, Cho CJ, Liu JR, et al. A streaming brain-to-voice neuroprosthesis to restore naturalistic communication. Nat Neurosci. 2025;28:902–912. doi:10.1038/s41593-025-01905-6.",
]


def add_endnotes_slide(items, part, total_parts):
    sl = prs.slides.add_slide(prs.slide_layouts[5])
    fill_title(sl, f"尾注 {part}/{total_parts}｜参考文献与证据状态", size=28)
    add_text(sl, "正文 [n] 与此处编号一一对应；预印本、单例和单臂研究保留原始限定。",
             0.72, 1.48, 11.7, 0.3, size=12, color=SUB, line_spacing=None)
    y = 1.88
    for ref in items:
        add_text(sl, ref, 0.72, y, 11.7, 0.68,
                 size=12, color=INK, font=FONT_EN, line_spacing=1.5)
        y += 0.69
    add_notes(sl, f"尾注第 {part} 页。正文数字、样本量和证据等级均通过编号指向这里。")


add_endnotes_slide(references[0:7], 1, 3)
add_endnotes_slide(references[7:14], 2, 3)
add_endnotes_slide(references[14:21], 3, 3)

# Add page numbers after all slides exist.
for idx, sl in enumerate(prs.slides, 1):
    if idx != 1:
        add_page_number(sl, idx, len(prs.slides))

prs.core_properties.title = "从解码到可用：脑机接口功能重建的证据链"
prs.core_properties.subject = "脑机接口文献组会汇报"
prs.core_properties.author = "组会汇报"

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
print(f"written: {OUT}")
print(f"slides: {len(prs.slides)}")
